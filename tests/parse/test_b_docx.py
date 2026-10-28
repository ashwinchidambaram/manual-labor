"""Tests 20 and 21: DOCX adapter against ground truth and virtual-page rules."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK

from manuallabor.contracts.enums import BlockType
from manuallabor.parse.docx import DocxResult, _parse_pages, build_docx, count_virtual_pages

FIX = Path(__file__).resolve().parents[1] / "fixtures"
TOK = re.compile(r"[^\s|•]+")


def _f1(a: str, b: str) -> float:
    ta, tb = TOK.findall(a), TOK.findall(b)
    if not ta and not tb:
        return 1.0
    pool = list(tb)
    hit = 0
    for t in ta:
        if t in pool:
            pool.remove(t)
            hit += 1
    if hit == 0:
        return 0.0
    p, r = hit / len(ta), hit / len(tb)
    return 2 * p * r / (p + r)


@pytest.fixture(scope="module")
def manual() -> DocxResult:
    return build_docx(FIX / "manual.docx", "manual")


@pytest.fixture(scope="module")
def gt() -> dict[str, Any]:
    return json.loads((FIX / "ground_truth" / "manual.json").read_text())


def test_20_blocks_match_ground_truth(manual: DocxResult, gt: dict[str, Any]) -> None:
    want = gt["blocks"]
    assert len(manual.blocks) == len(want) == 13
    for b, w in zip(manual.blocks, want, strict=True):
        assert (b.type.value, b.page) == (w["type"], w["page"])
        assert _f1(b.raw_text, w["text"]) >= 0.98
        assert b.bbox is None
        if w["type"] == "heading":
            assert b.heading_level == w["heading_level"]
            assert b.section_number == w["section_number"]
    assert [b.reading_order for b in manual.blocks] == list(range(13))


def test_20_figures_tables_and_parsers(manual: DocxResult) -> None:
    figs = [b.figure_id for b in manual.blocks if b.type == BlockType.figure]
    assert figs == ["FIG-manual-p001-01", "FIG-manual-p002-01"]
    assert [f.figure_id for f in manual.figures] == figs
    assert all(f.docx_image_rid and f.bbox is None for f in manual.figures)
    (table,) = [b for b in manual.blocks if b.type == BlockType.table]
    assert table.table is not None
    assert (table.table.n_rows, table.table.n_cols) == (4, 3)
    assert all(c.is_header for c in table.table.cells if c.row == 0)
    assert not any(c.is_header for c in table.table.cells if c.row > 0)
    assert set(table.table.shape_by_parser) == {"docx", "docling"}
    both = [b for b in manual.blocks if {"docx", "docling"} <= set(b.raw_text_by_parser)]
    assert len(both) / len(manual.blocks) >= 0.8
    assert {o.parser for o in manual.outputs} == {"docx", "docling"}
    assert manual.title == "Acme X200 Service Manual"


def test_20_report_covers_every_candidate(manual: DocxResult) -> None:
    ids_in_blocks = {b.id for b in manual.blocks}
    for d in manual.report.dispositions:
        if d.outcome != "skipped":
            assert d.block_id in ids_in_blocks
    keys = [(d.page, d.parser, d.order) for d in manual.report.dispositions]
    assert keys == sorted(keys) and len(keys) == len(set(keys))
    primary = [
        d for d in manual.report.dispositions if d.parser == "docx" and d.outcome == "primary"
    ]
    assert len(primary) == 13
    assert [p.page for p in manual.layout] == [1, 2]


def test_20_deterministic() -> None:
    a = build_docx(FIX / "manual.docx", "manual")
    b = build_docx(FIX / "manual.docx", "manual")
    assert [x.model_dump() for x in a.blocks] == [x.model_dump() for x in b.blocks]
    assert a.report.model_dump() == b.report.model_dump()


# ---------------------------------------------------------------- test 21 (unit)


def _page_of(path: Path) -> list[tuple[int, str]]:
    parsed = _parse_pages(path)
    return [(b.page, b.text) for p in parsed.pages for b in p.blocks]


def test_21_explicit_break_mid_paragraph(tmp_path: Path) -> None:
    d = Document()
    d.add_paragraph("first page text")
    p = d.add_paragraph()
    p.add_run("before the break")
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run("after the break")
    d.add_paragraph("tail")
    path = tmp_path / "break.docx"
    d.save(str(path))
    assert _page_of(path) == [
        (1, "first page text"),
        (1, "before the break"),
        (2, "after the break"),
        (2, "tail"),
    ]
    assert count_virtual_pages(path) == 2


def test_21_section_break(tmp_path: Path) -> None:
    d = Document()
    d.add_paragraph("one")
    d.add_section(WD_SECTION.NEW_PAGE)
    d.add_paragraph("two")
    d.add_section(WD_SECTION.CONTINUOUS)
    d.add_paragraph("three")
    path = tmp_path / "sect.docx"
    d.save(str(path))
    pages = _page_of(path)
    assert [t for _, t in pages] == ["one", "two", "three"]
    assert [p for p, _ in pages] == [1, 2, 3]
    assert count_virtual_pages(path) == 3


def test_21_999_block_rule(tmp_path: Path) -> None:
    d = Document()
    for i in range(1000):
        d.add_paragraph(f"para {i:04d}")
    path = tmp_path / "big.docx"
    d.save(str(path))
    pages = _page_of(path)
    assert [p for p, _ in pages[:999]] == [1] * 999
    assert pages[999] == (2, "para 0999")
    assert count_virtual_pages(path) == 2
    d2 = Document()
    for i in range(999):
        d2.add_paragraph(f"para {i:04d}")
    path2 = tmp_path / "exact.docx"
    d2.save(str(path2))
    assert count_virtual_pages(path2) == 1  # no trailing empty page


def test_21_lists_empty_paragraphs_and_merged_cells(tmp_path: Path) -> None:
    d = Document()
    d.add_paragraph("a", style="List Bullet")
    d.add_paragraph("")
    d.add_paragraph("b", style="List Bullet")
    d.add_paragraph("")
    d.add_paragraph("c")
    table = d.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 2).merge(table.cell(2, 2))
    table.cell(0, 0).text = "H"
    table.cell(1, 2).text = "V"
    path = tmp_path / "mix.docx"
    d.save(str(path))
    parsed = _parse_pages(path)
    blocks = parsed.pages[0].blocks
    assert [b.kind for b in blocks] == [BlockType.list, BlockType.paragraph, BlockType.table]
    assert blocks[0].text == "a\nb"
    assert len(parsed.pages[0].skipped) == 2
    tc = blocks[2].table
    assert tc is not None
    by = {(c.row, c.col): c for c in tc.cells}
    assert by[(0, 0)].col_span == 2 and by[(1, 2)].row_span == 2
    assert (0, 1) not in by and (2, 2) not in by
