"""Deterministic generator for the Manual Labor test fixtures (design 00-contracts section 8).

    uv run python -m tests.fixtures.gen.make_fixtures --out tests/fixtures [--check]

Ground truth is computed from the same coordinates that are drawn, never estimated afterwards.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import itertools
import math
import random
import re
import shutil
import struct
import sys
import tempfile
import zipfile
import zlib
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import numpy as np
import pymupdf as fitz
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import Table, TableStyle

from manuallabor.contracts.block import TableCell
from manuallabor.contracts.config import ModelRef, ProjectConfig, TaskModels
from manuallabor.contracts.enums import BlockType, DocKind, Extraction
from manuallabor.contracts.io import dump_json
from manuallabor.ids import doc_slug_from_filename, figure_id
from tests.fixtures.gen import content as C
from tests.fixtures.gen.ground_truth import (
    GroundTruth,
    GTBlock,
    GTFigure,
    GTSourceImage,
    GTTable,
    pixel_sha256,
)

SEED = 20261028
W, H = LETTER
LEFT, TOP, BOTTOM = 72.0, 72.0, 720.0
FULL_W = 468.0
GAP = 8.0
AUTHOR = "Manual Labor fixtures"
PDF_DATE = "D:20260101000000Z"

BODY = ("Helvetica", 10.0, 14.0)
H1 = ("Helvetica-Bold", 16.0, 20.0)
H2 = ("Helvetica-Bold", 12.0, 16.0)
CAP = ("Helvetica-Oblique", 9.0, 11.0)
HF = ("Helvetica", 9.0, 11.0)

FIXTURE_PDFS = ["clean", "twocol", "tables", "figures", "scanned"]
FIXTURE_FILES = [f"{s}.pdf" for s in FIXTURE_PDFS] + ["manual.docx"]
STEMS = FIXTURE_PDFS + ["manual"]


class LayoutOverflow(RuntimeError):
    pass


# ----------------------------------------------------------------------------- records


@dataclass
class Blk:
    page: int
    type: BlockType
    text: str
    raw_lines: list[str]
    bbox: list[float] | None
    heading_level: int | None = None
    section_number: str | None = None
    column: int | None = None
    is_hf: bool = False
    specs: list[str] = field(default_factory=list)
    hyphenated: bool = False
    table: GTTable | None = None
    index: int = -1

    @property
    def figure_refs(self) -> list[int]:
        return [int(n) for n in re.findall(r"see Figure (\d+)", self.text)]


@dataclass
class Fig:
    number: int | None
    page: int
    bbox: list[float] | None
    extraction: Extraction
    caption: str | None
    caption_blk: Blk | None
    column: int | None
    section: str | None
    source: GTSourceImage | None = None
    dup_of: Fig | None = None
    near: bool = False
    n: int = 0  # per-page order, assigned when the page is finished


# ----------------------------------------------------------------------------- text helpers


def sw(text: str, style: tuple[str, float, float]) -> float:
    return float(stringWidth(text, style[0], style[1]))


def wrap(
    text: str,
    style: tuple[str, float, float],
    width: float,
    hyphen: str | None = None,
) -> tuple[list[str], bool]:
    """Greedy wrap on spaces; optionally force a hyphenated break inside word `hyphen`."""
    lines: list[str] = []
    cur = ""
    done = False
    for word in text.split(" "):
        core = word.rstrip(",.;:)")
        punct = word[len(core) :]
        if hyphen and not done and core == hyphen and cur:
            head, tail = C.HYPHEN_WORDS[hyphen]
            trial = f"{cur} {head}-"
            if sw(f"{cur} {word}", style) > width and sw(trial, style) <= width:
                lines.append(trial)
                cur = tail + punct
                done = True
                continue
        trial = f"{cur} {word}" if cur else word
        if sw(trial, style) <= width:
            cur = trial
        else:
            lines.append(cur)
            cur = word
    lines.append(cur)
    return lines, done


def block_height(n_lines: int, style: tuple[str, float, float]) -> float:
    asc, desc = pdfmetrics.getAscentDescent(style[0], style[1])
    return asc - desc + (n_lines - 1) * style[2]


# ----------------------------------------------------------------------------- layout engine


class Layout:
    def __init__(
        self,
        path: Path,
        title: str,
        rng: random.Random,
        frames: list[tuple[float, float]] | None = None,
        hf_from: int = 1,
        keep: float = 60.0,
    ) -> None:
        self.rng = rng
        self.frames = frames or [(LEFT, FULL_W)]
        self.hf_from = hf_from
        self.keep = keep
        self.c = rl_canvas.Canvas(str(path), pagesize=LETTER, invariant=1, pageCompression=1)
        self.c.setTitle(title)
        self.c.setAuthor(AUTHOR)
        self.c.setSubject("Manual Labor test fixture")
        self.c.setCreator(AUTHOR)
        self.blocks: list[Blk] = []
        self.figs: list[Fig] = []
        self.images: dict[str, tuple[ImageReader, GTSourceImage]] = {}
        self.page = 0
        self.col = 0
        self.y = TOP
        self.section: str | None = None
        self._page_blocks: list[Blk] = []
        self._page_figs: list[Fig] = []
        self._start_page()

    # -- pages and frames
    def _col_tag(self) -> int | None:
        return self.col if len(self.frames) > 1 else None

    @property
    def fx(self) -> float:
        return self.frames[self.col][0]

    @property
    def fw(self) -> float:
        return self.frames[self.col][1]

    def _draw(
        self,
        x: float,
        baseline: float,
        lines: list[str],
        style: tuple[str, float, float],
    ) -> list[float]:
        font, size, leading = style
        self.c.setFillColorRGB(0, 0, 0)
        self.c.setFont(font, size)
        for i, line in enumerate(lines):
            self.c.drawString(x, baseline - i * leading, line)
        asc, desc = pdfmetrics.getAscentDescent(font, size)
        x1 = x + max(sw(line, style) for line in lines)
        last = baseline - (len(lines) - 1) * leading
        return [x, H - (baseline + asc), x1, H - (last + desc)]

    def _add(self, blk: Blk) -> Blk:
        self._page_blocks.append(blk)
        return blk

    def _start_page(self) -> None:
        self.page += 1
        self.col = 0
        self.y = TOP
        self._page_blocks = []
        self._page_figs = []
        if self.page >= self.hf_from:
            bbox = self._draw(LEFT, H - 45.0, [C.TITLE], HF)
            self._add(Blk(self.page, BlockType.header, C.TITLE, [C.TITLE], bbox, is_hf=True))

    def _finish_page(self) -> None:
        if self.page >= self.hf_from:
            label = f"Page {self.page}"
            bbox = self._draw(LEFT, 40.0, [label], HF)
            self._add(Blk(self.page, BlockType.footer, label, [label], bbox, is_hf=True))
        self.c.showPage()
        ordered = sorted(
            self._page_figs,
            key=lambda f: (f.column or 0, (f.bbox or [0, 0])[1], (f.bbox or [0])[0]),
        )
        for n, fig in enumerate(ordered, start=1):
            fig.n = n
        self.blocks.extend(self._page_blocks)
        self.figs.extend(self._page_figs)

    def page_break(self) -> None:
        self._finish_page()
        self._start_page()

    def next_frame(self) -> None:
        if self.col + 1 >= len(self.frames):
            raise LayoutOverflow(f"no next frame on page {self.page}")
        self.col += 1
        self.y = TOP

    def remaining(self) -> float:
        return BOTTOM - self.y

    def _need(self, h: float, what: str) -> None:
        if self.y + h > BOTTOM:
            raise LayoutOverflow(
                f"{what} (h={h:.0f}) does not fit on page {self.page} col {self.col} "
                f"(y={self.y:.0f})"
            )

    # -- text blocks
    def _text(
        self,
        btype: BlockType,
        text: str,
        lines: list[str],
        style: tuple[str, float, float],
        **kw: Any,
    ) -> Blk:
        h = block_height(len(lines), style)
        self._need(h, f"{btype.value} {text[:30]!r}")
        asc, _ = pdfmetrics.getAscentDescent(style[0], style[1])
        bbox = self._draw(self.fx, H - (self.y + asc), lines, style)
        self.y += h + GAP
        return self._add(Blk(self.page, btype, text, lines, bbox, column=self._col_tag(), **kw))

    def heading(self, text: str, level: int, style: tuple[str, float, float] | None = None) -> Blk:
        style = style or (H1 if level == 1 else H2)
        lines, _ = wrap(text, style, self.fw)
        self._need(block_height(len(lines), style) + self.keep, f"heading {text!r}")
        m = re.match(r"^(\d+(?:\.\d+)*)\.?\s", text)
        sec = m.group(1) if m else None
        if sec:
            self.section = sec
        return self._text(
            BlockType.heading, text, lines, style, heading_level=level, section_number=sec
        )

    def para(
        self,
        text: str,
        specs: list[str] | None = None,
        hyphen: str | None = None,
        style: tuple[str, float, float] = BODY,
    ) -> Blk:
        lines, done = wrap(text, style, self.fw, hyphen)
        if hyphen and not done:
            raise LayoutOverflow("hyphenation not realised")
        return self._text(
            BlockType.paragraph, text, lines, style, specs=list(specs or []), hyphenated=done
        )

    def para_auto(
        self,
        specs: list[str] | None = None,
        fig_refs: list[int] | None = None,
        table_refs: list[int] | None = None,
        n: int | None = None,
    ) -> Blk:
        specs = specs if specs is not None else C.random_specs(self.rng)
        text, used = C.make_paragraph(
            self.rng, n or self.rng.randint(3, 5), specs, fig_refs, table_refs
        )
        return self.para(text, used)

    def para_hyph(self, specs: list[str] | None = None) -> Blk:
        for _ in range(400):
            word = self.rng.choice(sorted(C.HYPHEN_WORDS))
            sp = specs if specs is not None else C.random_specs(self.rng)
            text, used = C.make_paragraph(self.rng, self.rng.randint(4, 6), sp, hyphen_word=word)
            lines, done = wrap(text, BODY, self.fw, word)
            if done:
                return self.para(text, used, hyphen=word)
        raise LayoutOverflow("could not force a hyphenated line break")

    def fill(self, max_n: int = 30) -> int:
        """Add paragraphs until the next one no longer fits in the current frame."""
        placed = 0
        while placed < max_n:
            specs = C.random_specs(self.rng)
            text, used = C.make_paragraph(self.rng, self.rng.randint(3, 6), specs)
            lines, _ = wrap(text, BODY, self.fw)
            if self.y + block_height(len(lines), BODY) > BOTTOM:
                break
            self._text(BlockType.paragraph, text, lines, BODY, specs=used)
            placed += 1
        return placed

    def bullets(self, items: list[str]) -> Blk:
        lines = [f"• {it}" for it in items]
        for ln in lines:
            if sw(ln, BODY) > self.fw:
                raise LayoutOverflow(f"bullet too long for frame: {ln!r}")
        specs = [s for s in C.SPECS if any(s in it for it in items)]
        return self._text(BlockType.list, "\n".join(items), lines, BODY, specs=specs)

    # -- figures
    def add_image(self, key: str, img: Image.Image) -> GTSourceImage:
        info = GTSourceImage(sha256=pixel_sha256(img), width=img.width, height=img.height)
        self.images[key] = (ImageReader(img), info)
        return info

    def _caption_lines(self, number: int | None, text: str | None) -> tuple[str, list[str]] | None:
        if number is None or text is None:
            return None
        cap = f"Figure {number}. {text}"
        return cap, wrap(cap, CAP, self.fw)[0]

    def _figure(
        self,
        w: float,
        h: float,
        number: int | None,
        caption_text: str | None,
        extraction: Extraction,
        draw: Any,
        source: GTSourceImage | None = None,
    ) -> Fig:
        cap = self._caption_lines(number, caption_text)
        total = h + (4 + block_height(len(cap[1]), CAP) if cap else 0)
        self._need(total, f"figure {number}")
        x, y = self.fx, self.y
        draw(x, H - (y + h), w, h)
        bbox = [x, y, x + w, y + h]
        fblk = self._add(Blk(self.page, BlockType.figure, "", [], bbox, column=self._col_tag()))
        fblk.index = -1
        self.y += h
        cap_blk = None
        if cap:
            self.y += 4
            cap_blk = self._text(BlockType.caption, cap[0], cap[1], CAP)
        else:
            self.y += GAP
        fig = Fig(
            number,
            self.page,
            bbox,
            extraction,
            cap[0] if cap else None,
            cap_blk,
            self._col_tag(),
            self.section,
            source,
        )
        self._page_figs.append(fig)
        return fig

    def raster(
        self,
        key: str,
        w: float,
        h: float,
        number: int | None,
        caption: str | None,
        separate_object: bool = False,
    ) -> Fig:
        """separate_object=True makes reportlab write a second PDF image object (own xref)
        with identical pixel data: its cache name hashes (pixels + str(mask)), so passing
        mask="auto" (a no-op for an image without alpha) changes the name."""
        reader, info = self.images[key]

        def draw(x: float, ypdf: float, ww: float, hh: float) -> None:
            self.c.drawImage(reader, x, ypdf, ww, hh, mask="auto" if separate_object else None)

        return self._figure(w, h, number, caption, Extraction.native, draw, info)

    def vector(self, w: float, h: float, number: int | None, caption: str | None) -> Fig:
        rng = self.rng
        colors = [(0.1, 0.1, 0.5), (0.6, 0.1, 0.1), (0.1, 0.45, 0.2), (0.3, 0.3, 0.3)]

        def draw(x: float, ypdf: float, ww: float, hh: float) -> None:
            c = self.c
            c.saveState()
            # clip every primitive to the figure frame; geometry below also stays inside it
            clip = c.beginPath()
            clip.rect(x, ypdf, ww, hh)
            c.clipPath(clip, stroke=0, fill=0)
            c.setLineWidth(0.5)
            c.setStrokeColorRGB(0.5, 0.5, 0.5)
            c.rect(x, ypdf, ww, hh, stroke=1, fill=0)
            m = 6.0
            for _ in range(rng.randint(8, 12)):
                c.setStrokeColorRGB(*rng.choice(colors))
                c.setLineWidth(rng.choice([0.8, 1.2, 2.0]))
                kind = rng.choice(["line", "rect", "circle", "bezier"])
                px = [x + m + rng.random() * (ww - 2 * m) for _ in range(4)]
                py = [ypdf + m + rng.random() * (hh - 2 * m) for _ in range(4)]
                if kind == "line":
                    c.line(px[0], py[0], px[1], py[1])
                elif kind == "rect":
                    rw, rh = abs(px[0] - px[1]) * 0.6, abs(py[0] - py[1]) * 0.6
                    rx = min(min(px[:2]), x + ww - m - rw)
                    ry = min(min(py[:2]), ypdf + hh - m - rh)
                    c.rect(rx, ry, rw, rh)
                elif kind == "circle":
                    r = min(4 + rng.random() * min(ww, hh) / 8, min(ww, hh) / 2 - m - 1)
                    cx = min(max(px[0], x + m + r), x + ww - m - r)
                    cy = min(max(py[0], ypdf + m + r), ypdf + hh - m - r)
                    c.circle(cx, cy, r)
                else:
                    p = c.beginPath()
                    p.moveTo(px[0], py[0])
                    p.curveTo(px[1], py[1], px[2], py[2], px[3], py[3])
                    c.drawPath(p, stroke=1, fill=0)
            c.restoreState()

        return self._figure(w, h, number, caption, Extraction.vector, draw)

    def composite(
        self, key: str, w: float, h: float, number: int | None, caption: str | None
    ) -> Fig:
        reader, info = self.images[key]

        def draw(x: float, ypdf: float, ww: float, hh: float) -> None:
            c = self.c
            c.drawImage(reader, x, ypdf, ww, hh)
            c.saveState()
            c.setStrokeColorRGB(0.85, 0.1, 0.1)
            c.setFillColorRGB(0.85, 0.1, 0.1)
            c.setLineWidth(1.5)
            for i, (fx, fy, tx, ty, lab) in enumerate(
                [
                    (0.15, 0.2, 0.4, 0.5, "A"),
                    (0.85, 0.25, 0.6, 0.55, "B"),
                    (0.5, 0.9, 0.5, 0.65, "C"),
                ]
            ):
                x0, y0, x1, y1 = x + fx * ww, ypdf + fy * hh, x + tx * ww, ypdf + ty * hh
                c.line(x0, y0, x1, y1)
                p = c.beginPath()
                p.moveTo(x1, y1)
                dx, dy = x0 - x1, y0 - y1
                norm = (dx * dx + dy * dy) ** 0.5
                ux, uy = dx / norm, dy / norm
                p.lineTo(x1 + 8 * ux - 4 * uy, y1 + 8 * uy + 4 * ux)
                p.lineTo(x1 + 8 * ux + 4 * uy, y1 + 8 * uy - 4 * ux)
                p.close()
                c.drawPath(p, stroke=0, fill=1)
                c.setFont("Helvetica-Bold", 10)
                c.drawString(x0 - 3, y0 + (4 if i < 2 else -12), lab)
            c.restoreState()

        return self._figure(w, h, number, caption, Extraction.crop, draw, info)

    # -- tables
    def table(self, spec: dict[str, Any], number: int) -> Blk:
        rows: list[list[str | None]] = spec["rows"]
        n_rows, n_cols = len(rows), len(rows[0])
        widths: list[float] = spec["col_widths"]
        hr: int = spec["header_rows"]
        data = [["" if v is None else v for v in row] for row in rows]
        cells: list[TableCell] = []
        span_at = {(r, c): (rs, cs) for r, c, rs, cs in spec["spans"]}
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                if v is None:
                    continue
                rs, cs = span_at.get((r, c), (1, 1))
                if sw(v, ("Helvetica-Bold", 9.0, 0)) + 12 > sum(widths[c : c + cs]):
                    raise LayoutOverflow(f"table cell too narrow: {v!r}")
                cells.append(
                    TableCell(row=r, col=c, row_span=rs, col_span=cs, text=v, is_header=r < hr)
                )
        style = [
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("FONTNAME", (0, 0), (-1, hr - 1), "Helvetica-Bold"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ]
        if spec["bordered"]:
            style += [
                ("GRID", (0, 0), (-1, -1), 0.5, (0, 0, 0)),
                ("BACKGROUND", (0, 0), (-1, hr - 1), (0.88, 0.88, 0.88)),
            ]
        for (r, c), (rs, cs) in span_at.items():
            style.append(("SPAN", (c, r), (c + cs - 1, r + rs - 1)))
        tbl = Table(data, colWidths=widths, rowHeights=[18] * n_rows)
        tbl.setStyle(TableStyle(style))
        tw, th = tbl.wrapOn(self.c, W, H)
        cap_text: str = spec["caption"]
        cap_lines = wrap(cap_text, CAP, self.fw)[0]
        self._need(block_height(len(cap_lines), CAP) + 4 + th, f"table {number}")
        self._text(BlockType.caption, cap_text, cap_lines, CAP)
        self.y += 4 - GAP
        x, y = self.fx, self.y
        tbl.drawOn(self.c, x, H - (y + th))
        raw = []
        for r in range(n_rows):
            parts = [t.text for t in cells if t.row == r and t.text]
            raw.append(" ".join(parts))
        text = "\n".join(
            " | ".join(t.text for t in cells if t.row == r and t.text) for r in range(n_rows)
        )
        all_text = " ".join(t.text for t in cells)
        gt = GTTable(
            n_rows=n_rows,
            n_cols=n_cols,
            cells=sorted(cells, key=lambda t: (t.row, t.col)),
            bordered=bool(spec["bordered"]),
            has_merged=any(rs > 1 or cs > 1 for rs, cs in span_at.values()),
            ragged=bool(spec.get("ragged", False)),
            caption=cap_text,
        )
        blk = self._add(
            Blk(
                self.page,
                BlockType.table,
                text,
                raw,
                [x, y, x + tw, y + th],
                column=self._col_tag(),
                specs=[s for s in C.SPECS if s in all_text],
                table=gt,
            )
        )
        self.y += th + GAP
        return blk

    def close(self) -> int:
        self._finish_page()
        self.c.save()
        for i, blk in enumerate(self.blocks):
            blk.index = i
        return self.page


# ----------------------------------------------------------------------------- images


def png_bytes(img: Image.Image) -> bytes:
    """Encode an image as an 8-bit RGB PNG with Python's own ``zlib``, byte-stable everywhere.

    Pillow's PNG writer goes through its bundled zlib-ng, which emits different deflate
    streams on macOS arm64 and Linux x86-64 for identical pixels, so any PNG whose bytes end
    up inside a fixture is written here instead: signature, IHDR, one IDAT (filter type 0 on
    every scanline, ``zlib.compress(level=9)``) and IEND, each chunk with ``zlib.crc32``.
    """
    arr = np.asarray(img.convert("RGB"), dtype=np.uint8)
    h, w = arr.shape[:2]
    rows = np.zeros((h, 1 + w * 3), dtype=np.uint8)
    rows[:, 1:] = arr.reshape(h, w * 3)

    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))

    ihdr = struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(rows.tobytes(), 9))
        + chunk(b"IEND", b"")
    )


def box_downscale(img: Image.Image, size: tuple[int, int]) -> Image.Image:
    """Area-average (box filter) downscale by an integer factor per axis, in integer arithmetic.

    Replaces ``Image.resize(..., Image.LANCZOS)``, whose floating-point kernel gave one
    different pixel on Linux x86-64 than on macOS arm64. Each output pixel is the mean of its
    ``fx * fy`` source pixels, summed in int64 and rounded half up with integer division, so
    the result is identical on every platform.
    """
    ow, oh = size
    rgb = np.asarray(img.convert("RGB"), dtype=np.int64)
    h, w = rgb.shape[:2]
    if w % ow or h % oh:
        raise ValueError(f"box_downscale needs integer factors: {w}x{h} -> {ow}x{oh}")
    fx, fy = w // ow, h // oh
    n = fx * fy
    sums = rgb.reshape(oh, fy, ow, fx, 3).sum(axis=(1, 3))
    return Image.fromarray(((sums + n // 2) // n).astype(np.uint8), "RGB")


def bresenham(x0: int, y0: int, x1: int, y1: int) -> tuple[np.ndarray, np.ndarray]:
    """Integer Bresenham points from (x0, y0) to (x1, y1), both ends included."""
    dx, sx = abs(x1 - x0), (1 if x0 < x1 else -1)
    dy, sy = -abs(y1 - y0), (1 if y0 < y1 else -1)
    err = dx + dy
    xs: list[int] = []
    ys: list[int] = []
    x, y = x0, y0
    while True:
        xs.append(x)
        ys.append(y)
        if x == x1 and y == y1:
            break
        e2 = 2 * err
        if e2 >= dy:
            err += dy
            x += sx
        if e2 <= dx:
            err += dx
            y += sy
    return np.array(xs, dtype=np.int64), np.array(ys, dtype=np.int64)


def wide_line(
    img: Image.Image, xy: tuple[int, int, int, int], fill: tuple[int, int, int], width: int
) -> None:
    """Draw a ``width``-pixel line in integer arithmetic only, in place.

    Replaces ``ImageDraw.line(width > 1)``: Pillow's ``ImagingDrawWideLine`` builds a polygon
    from ``hypot`` and float edge slopes, and the float rounding differs between macOS arm64
    and Linux x86-64 (one pixel flipped in a figures.pdf image). Here the centre line is
    Bresenham, and at every centre point a span is filled across the minor axis (vertical for
    an x-major line, horizontal for a y-major one). The span is ``width * length / major``
    pixels, rounded, with ``length = isqrt(dx² + dy²)``, so the perpendicular thickness stays
    close to ``width`` as with Pillow. A square pen was tried and rejected: it fattens the
    diagonals so much that the twocol composite's phash came within the near-duplicate
    threshold (6 <= 8) of its base image. Rectangle and ellipse outlines stay on ``ImageDraw``:
    Pillow draws those with integer hlines and integer quarter-ellipse steps.
    """
    x0, y0, x1, y1 = xy
    xs, ys = bresenham(x0, y0, x1, y1)
    dx, dy = abs(x1 - x0), abs(y1 - y0)
    major = max(dx, dy, 1)
    span = max(width, (width * math.isqrt(dx * dx + dy * dy) + major // 2) // major)
    offs = np.arange(-(span // 2), span - span // 2, dtype=np.int64)
    if dx >= dy:
        px = np.repeat(xs, len(offs))
        py = (ys[:, None] + offs[None, :]).ravel()
    else:
        px = (xs[:, None] + offs[None, :]).ravel()
        py = np.repeat(ys, len(offs))
    keep = (px >= 0) & (px < img.width) & (py >= 0) & (py < img.height)
    arr = np.array(img)
    arr[py[keep], px[keep]] = fill
    img.paste(Image.fromarray(arr, "RGB"))


def line_art(rng: random.Random, w: int, h: int) -> Image.Image:
    img = Image.new("RGB", (w, h), (255, 255, 255))
    d = ImageDraw.Draw(img)
    palette = [(20, 20, 120), (150, 30, 30), (30, 110, 50), (60, 60, 60), (200, 120, 0)]
    lw = max(2, w // 250)
    d.rectangle([2, 2, w - 3, h - 3], outline=(0, 0, 0), width=lw)
    for _ in range(rng.randint(16, 24)):
        kind = rng.choice(["line", "rect", "ellipse"])
        x0, x1 = sorted((rng.randint(10, w - 10), rng.randint(10, w - 10)))
        y0, y1 = sorted((rng.randint(10, h - 10), rng.randint(10, h - 10)))
        col = rng.choice(palette)
        if kind == "line":
            wide_line(img, (x0, y0, x1, y1), col, lw)
        elif kind == "rect":
            d.rectangle([x0, y0, x1, y1], outline=col, width=lw)
        else:
            d.ellipse([x0, y0, x1 + 4, y1 + 4], outline=col, width=lw)
    return img


# ----------------------------------------------------------------------------- fixtures


def build_clean(path: Path) -> tuple[Layout, int]:
    rng = random.Random(SEED)
    lay = Layout(path, C.TITLE, rng, hf_from=2, keep=60)
    # page 1: title page
    lay.y = 250.0
    lay.heading(C.TITLE, 1, style=("Helvetica-Bold", 28.0, 34.0))
    lay.para(C.SUBTITLE, style=("Helvetica", 14.0, 18.0))
    imgs = [line_art(rng, 900, 600) for _ in range(3)]
    for i, img in enumerate(imgs, start=1):
        lay.add_image(f"F{i}", img)
    bl = itertools.cycle(C.BULLET_LISTS)
    sec = C.SECTIONS_CLEAN

    def page(h1: str | None, h2: str | None, hyph: bool, bullets: bool, fill: bool = True) -> None:
        lay.page_break()
        if h1:
            lay.heading(h1, 1)
        if h2:
            lay.heading(h2, 2)
        lay.para_auto(specs=[C.SPECS[0], C.SPECS[1]] if lay.page == 2 else None)
        if hyph:
            lay.para_hyph()
        if bullets:
            lay.bullets(next(bl))
        if fill:
            lay.fill()

    page(sec[0][0], sec[0][1][0], True, True)  # p2
    page(None, sec[0][1][1], False, True)  # p3
    page(sec[1][0], sec[1][1][0], True, True)  # p4
    # p5: 2.2 with Figure 3 (page before the other two figures)
    lay.page_break()
    lay.heading(sec[1][1][1], 2)
    lay.para_auto(specs=[C.SPECS[2], C.SPECS[3]])
    lay.para_hyph()
    lay.bullets(next(bl))
    lay.para_auto(n=3)
    lay.raster("F3", 300, 200, 3, C.FIGURE_CAPTIONS[2])
    # p6: 2.3 with Figures 1 and 2
    lay.page_break()
    lay.heading(sec[1][1][2], 2)
    lay.para_auto(specs=[C.SPECS[5]], fig_refs=[1, 2, 3], n=3)
    lay.raster("F1", 300, 200, 1, C.FIGURE_CAPTIONS[0])
    lay.raster("F2", 300, 200, 2, C.FIGURE_CAPTIONS[1])
    lay.fill()
    page(sec[2][0], sec[2][1][0], True, True)  # p7
    page(None, sec[2][1][1], False, True)  # p8
    page(sec[3][0], sec[3][1][0], False, True)  # p9
    page(None, sec[3][1][1], False, False)  # p10
    # p11: 4.3 + section 5 + 5.1
    lay.page_break()
    lay.heading(sec[3][1][2], 2)
    lay.para_auto(n=3)
    lay.bullets(next(bl))
    lay.heading(sec[4][0], 1)
    lay.heading(sec[4][1][0], 2)
    lay.para_auto(n=4)
    page(None, sec[4][1][1], False, False)  # p12
    return lay, lay.close()


def build_twocol(path: Path) -> tuple[Layout, int]:
    rng = random.Random(SEED)
    lay = Layout(path, C.TITLE, rng, frames=[(72.0, 225.0), (315.0, 225.0)], keep=170)
    for i in (1, 2, 3):
        lay.add_image(f"T{i}", line_art(rng, 900, 600))
    titles = [
        ("1. Cooling System", "1.1 Overview", "1.2 Components"),
        ("2. Lubrication", "2.1 Oil Circuit", "2.2 Filters"),
        ("3. Electrical Layout", "3.1 Harness", "3.2 Connectors"),
        ("4. Drive Train", "4.1 Belt Path", "4.2 Pulleys"),
        ("5. Control Board", "5.1 Inputs", "5.2 Outputs"),
        ("6. Storage", "6.1 Preparation", "6.2 Restart"),
    ]
    for p, (t1, t2, t3) in enumerate(titles, start=1):
        if p > 1:
            lay.page_break()
        lay.heading(t1, 1)
        lay.heading(t2, 2)
        if p == 2:
            lay.para_auto(specs=[C.SPECS[0]], fig_refs=[1, 2])
            lay.raster("T1", 225, 150, 1, C.FIGURE_CAPTIONS[0])
        if p == 4:
            lay.para_auto(specs=[C.SPECS[1]], fig_refs=[3])
            lay.composite("T3", 225, 150, 3, C.FIGURE_CAPTIONS[3])
        lay.fill()
        lay.next_frame()
        if p == 2:
            lay.raster("T2", 225, 150, 2, C.FIGURE_CAPTIONS[1])
        lay.heading(t3, 2)
        lay.fill()
    return lay, lay.close()


def build_tables(path: Path) -> tuple[Layout, int]:
    rng = random.Random(SEED)
    lay = Layout(path, C.TITLE, rng, keep=60)
    heads = [
        "1. Torque Data",
        "2. Electrical Data",
        "3. Parts List",
        "4. Service Intervals",
        "5. Connector Data",
        "6. Fluid Data",
    ]
    for i, (spec, head) in enumerate(zip(C.TABLES, heads, strict=True), start=1):
        if i > 1:
            lay.page_break()
        lay.heading(head, 1)
        lay.para_auto(specs=[C.SPECS[i - 1]], table_refs=[i])
        lay.table(spec, i)
        lay.fill()
    return lay, lay.close()


def build_figures(path: Path) -> tuple[Layout, int]:
    rng = random.Random(SEED)
    lay = Layout(path, C.TITLE, rng, keep=60)
    sizes = {
        "R1": (1200, 800),
        "R2": (600, 400),
        "R3": (900, 900),
        "R4": (1600, 1000),
        "R5": (600, 400),
        "R6": (1200, 800),
    }
    imgs = {k: line_art(rng, *wh) for k, wh in sizes.items()}
    for k, img in imgs.items():
        lay.add_image(k, img)
    lay.add_image("R1n", box_downscale(imgs["R1"], (600, 400)))
    # (kind, key, w, h, flags)
    plan: list[list[tuple[str, str, float, float, str]]] = [
        [("v", "V1", 260, 160, ""), ("r", "R2", 240, 160, "")],
        [("r", "R1", 240, 160, ""), ("v", "V2", 260, 160, "")],
        [("r", "R3", 200, 200, "nocap"), ("v", "V3", 240, 150, "")],
        [("r", "R4", 320, 200, "xref")],
        [("v", "V4", 260, 160, ""), ("r", "R5", 240, 160, "noref")],
        [("r", "R1", 240, 160, "dup"), ("v", "V5", 260, 160, "")],
        [("r", "R1n", 240, 160, "near"), ("r", "R6", 260, 170, "")],
        [("v", "V6", 300, 200, "xref")],
    ]
    number = 0
    nums: dict[tuple[int, int], int | None] = {}
    for pi, items in enumerate(plan):
        for j, it in enumerate(items):
            if it[4] == "nocap":
                nums[(pi, j)] = None
            else:
                number += 1
                nums[(pi, j)] = number
    heads = ["Drive", "Idler", "Cover", "Board", "Coolant", "Valve", "Harness", "Fan"]
    first: Fig | None = None
    for pi, items in enumerate(plan):
        if pi:
            lay.page_break()
        lay.heading(f"{pi + 1}. {heads[pi]} Illustrations", 1)
        refs = [
            nums[(pi, j)]
            for j, it in enumerate(items)
            if nums[(pi, j)] is not None and it[4] not in ("xref", "noref")
        ]
        if pi + 1 < len(plan):
            refs += [nums[(pi + 1, j)] for j, it in enumerate(plan[pi + 1]) if it[4] == "xref"]
        lay.para_auto(fig_refs=[r for r in refs if r is not None] or None)
        for j, (kind, key, w, h, flag) in enumerate(items):
            n = nums[(pi, j)]
            cap = C.FIGURE_CAPTIONS[(n or 0) % len(C.FIGURE_CAPTIONS)] if n else None
            fig = (
                lay.raster(key, w, h, n, cap, separate_object=flag == "dup")
                if kind == "r"
                else lay.vector(w, h, n, cap)
            )
            if key == "R1" and flag == "":
                first = fig
            if flag == "dup":
                fig.dup_of = first
            if flag == "near":
                fig.dup_of = first
                fig.near = True
        lay.fill()
    return lay, lay.close()


def build_scanned(clean_pdf: Path, path: Path) -> int:
    rng = np.random.default_rng(SEED)
    src = fitz.open(str(clean_pdf))
    out = fitz.open()
    for page in src:
        pix = page.get_pixmap(dpi=300, colorspace=fitz.csGRAY, alpha=False)
        img = Image.frombytes("L", (pix.width, pix.height), pix.samples)
        img = img.rotate(0.5, resample=Image.BICUBIC, fillcolor=255)
        arr = np.asarray(img, dtype=np.float32)
        arr = arr + rng.normal(0.0, 6.0, size=arr.shape).astype(np.float32)
        img = Image.fromarray(np.clip(np.rint(arr), 0, 255).astype(np.uint8), mode="L")
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=80, optimize=False, subsampling=0)
        p = out.new_page(width=W, height=H)
        p.insert_image(p.rect, stream=buf.getvalue())
    out.set_metadata(
        {
            "title": C.TITLE,
            "author": AUTHOR,
            "subject": "Manual Labor test fixture (scanned)",
            "creator": AUTHOR,
            "producer": AUTHOR,
            "keywords": "",
            "creationDate": PDF_DATE,
            "modDate": PDF_DATE,
        }
    )
    n = out.page_count
    out.save(str(path), garbage=4, deflate=True, no_new_id=True)
    out.close()
    src.close()
    return n


def normalize_zip(path: Path) -> None:
    with zipfile.ZipFile(path) as zin:
        entries = [(info.filename, zin.read(info.filename)) for info in zin.infolist()]
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries:
            zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            zout.writestr(zi, data)
    tmp.replace(path)


def build_docx(path: Path) -> tuple[list[Blk], list[Fig]]:
    import datetime as dt

    import docx
    from docx.shared import Inches

    rng = random.Random(SEED)
    d = docx.Document()
    cp = d.core_properties
    stamp = dt.datetime(2026, 1, 1, 0, 0, 0)
    cp.created = stamp
    cp.modified = stamp
    cp.last_modified_by = AUTHOR
    cp.author = AUTHOR
    cp.title = C.TITLE
    cp.revision = 1
    cp.last_printed = stamp
    blocks: list[Blk] = []
    figs: list[Fig] = []
    state = {"page": 1, "section": None}

    def add(btype: BlockType, text: str, raw: list[str], **kw: Any) -> Blk:
        b = Blk(int(state["page"]), btype, text, raw, None, **kw)  # type: ignore[arg-type]
        b.index = len(blocks)
        blocks.append(b)
        return b

    def heading(text: str, level: int) -> None:
        d.add_heading(text, level=level)
        m = re.match(r"^(\d+(?:\.\d+)*)\.?\s", text)
        state["section"] = m.group(1) if m else None
        add(BlockType.heading, text, [text], heading_level=level, section_number=state["section"])

    def para(text: str, specs: list[str]) -> None:
        d.add_paragraph(text)
        add(BlockType.paragraph, text, [text], specs=specs)

    def picture(key: int, number: int, w_in: float) -> None:
        img = line_art(rng, 900, 600)
        d.add_picture(io.BytesIO(png_bytes(img)), width=Inches(w_in))
        add(BlockType.figure, "", [])
        cap = f"Figure {number}. {C.FIGURE_CAPTIONS[key]}"
        d.add_paragraph(cap, style="Caption")
        cblk = add(BlockType.caption, cap, [cap])
        figs.append(
            Fig(
                number,
                int(state["page"]),
                None,
                Extraction.native,
                cap,
                cblk,
                None,
                state["section"],  # type: ignore[arg-type]
                GTSourceImage(sha256=pixel_sha256(img), width=900, height=600),
            )
        )

    heading("1. Overview", 1)
    t, s = C.make_paragraph(rng, 3, [C.SPECS[0]])
    para(t, s)
    heading("1.1 Required Items", 2)
    t, s = C.make_paragraph(rng, 3, [C.SPECS[3]], fig_refs=[1])
    para(t, s)
    items = C.BULLET_LISTS[0]
    for it in items:
        d.add_paragraph(it, style="List Bullet")
    add(
        BlockType.list,
        "\n".join(items),
        [f"• {it}" for it in items],
        specs=[s_ for s_ in C.SPECS if any(s_ in it for it in items)],
    )
    rows = [
        ["Item", "Size", "Torque"],
        ["Cover bolt", "M8", "35 N·m"],
        ["Bracket bolt", "M10", "48 N·m"],
        ["Fan bolt", "M5", "6 N·m"],
    ]
    tbl = d.add_table(rows=4, cols=3)
    tbl.style = "Table Grid"
    cells = []
    for r, row in enumerate(rows):
        for c_, v in enumerate(row):
            tbl.cell(r, c_).text = v
            cells.append(TableCell(row=r, col=c_, text=v, is_header=r == 0))
    all_text = " ".join(r_ for row in rows for r_ in row)
    add(
        BlockType.table,
        "\n".join(" | ".join(r_) for r_ in rows),
        [" ".join(r_) for r_ in rows],
        specs=[s_ for s_ in C.SPECS if s_ in all_text],
        table=GTTable(
            n_rows=4,
            n_cols=3,
            cells=cells,
            bordered=True,
            has_merged=False,
            ragged=False,
            caption=None,
        ),
    )
    picture(0, 1, 3.0)
    d.add_page_break()
    state["page"] = 2
    heading("2. Maintenance", 1)
    t, s = C.make_paragraph(rng, 4, [C.SPECS[2]], fig_refs=[2])
    para(t, s)
    picture(1, 2, 3.0)
    t, s = C.make_paragraph(rng, 3, [C.SPECS[6]])
    para(t, s)
    d.save(str(path))
    normalize_zip(path)
    return blocks, figs


# ----------------------------------------------------------------------------- ground truth


def build_gt(
    fixture: str,
    filename: str,
    kind: DocKind,
    sha256: str,
    page_count: int,
    blocks: list[Blk],
    figs: list[Fig],
    *,
    paged: bool = True,
    header: bool = True,
    derived_from: str | None = None,
    text_is_reference: bool = False,
    scanned: bool = False,
) -> GroundTruth:
    slug = doc_slug_from_filename(filename)

    def fid(f: Fig) -> str:
        return figure_id(slug, f.page, f.n)

    gblocks = []
    for b in blocks:
        gblocks.append(
            GTBlock(
                index=b.index,
                page=b.page,
                type=b.type,
                text=b.text,
                raw_lines=b.raw_lines,
                bbox=b.bbox,
                heading_level=b.heading_level,
                section_number=b.section_number,
                column=b.column,
                is_header_footer=b.is_hf,
                numeric_tokens=re.findall(r"\d+(?:[.,]\d+)*", b.text),
                numeric_specs=b.specs,
                has_numeric_spec=bool(b.specs),
                hyphenated=b.hyphenated,
                figure_refs=b.figure_refs,
                table=b.table,
            )
        )
    gfigs = []
    for f in figs:
        refs = (
            [b.index for b in blocks if f.number in b.figure_refs] if f.number is not None else []
        )
        gfigs.append(
            GTFigure(
                number=f.number,
                expected_id=fid(f),
                page=f.page,
                bbox=f.bbox,
                expected_extraction=Extraction.crop if scanned else f.extraction,
                caption=f.caption,
                caption_block_index=f.caption_blk.index if f.caption_blk else None,
                referenced_by=refs,
                duplicate_of=None if scanned or f.dup_of is None else fid(f.dup_of),
                near_duplicate=False if scanned else f.near,
                source_image=None if scanned else f.source,
                section_number=f.section,
            )
        )
    return GroundTruth(
        fixture=fixture,
        file=filename,
        sha256=sha256,
        kind=kind,
        page_count=page_count,
        page_size=[612.0, 792.0] if paged else None,
        expected_slug=slug,
        title=C.TITLE,
        product_model=C.PRODUCT_MODEL,
        header_text=C.TITLE if header else None,
        footer_template="Page {n}" if header else None,
        derived_from=derived_from,
        text_is_reference=text_is_reference,
        blocks=gblocks,
        figures=gfigs,
    )


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(text.encode("utf-8"))


def generate(out: Path) -> dict[str, str]:
    out.mkdir(parents=True, exist_ok=True)
    gt_dir = out / "ground_truth"
    gt_dir.mkdir(exist_ok=True)
    gts: dict[str, GroundTruth] = {}

    clean_lay, n = build_clean(out / "clean.pdf")
    assert n == 12, f"clean.pdf has {n} pages"
    gts["clean"] = build_gt(
        "clean",
        "clean.pdf",
        DocKind.born_digital_pdf,
        sha(out / "clean.pdf"),
        n,
        clean_lay.blocks,
        clean_lay.figs,
    )
    lay, n = build_twocol(out / "twocol.pdf")
    assert n == 6, f"twocol.pdf has {n} pages"
    gts["twocol"] = build_gt(
        "twocol",
        "twocol.pdf",
        DocKind.born_digital_pdf,
        sha(out / "twocol.pdf"),
        n,
        lay.blocks,
        lay.figs,
    )
    lay, n = build_tables(out / "tables.pdf")
    assert n == 6, f"tables.pdf has {n} pages"
    gts["tables"] = build_gt(
        "tables",
        "tables.pdf",
        DocKind.born_digital_pdf,
        sha(out / "tables.pdf"),
        n,
        lay.blocks,
        lay.figs,
    )
    lay, n = build_figures(out / "figures.pdf")
    assert n == 8, f"figures.pdf has {n} pages"
    gts["figures"] = build_gt(
        "figures",
        "figures.pdf",
        DocKind.born_digital_pdf,
        sha(out / "figures.pdf"),
        n,
        lay.blocks,
        lay.figs,
    )
    n = build_scanned(out / "clean.pdf", out / "scanned.pdf")
    gts["scanned"] = build_gt(
        "scanned",
        "scanned.pdf",
        DocKind.scanned_pdf,
        sha(out / "scanned.pdf"),
        n,
        clean_lay.blocks,
        clean_lay.figs,
        derived_from="clean",
        text_is_reference=True,
        scanned=True,
    )
    dblocks, dfigs = build_docx(out / "manual.docx")
    for f in dfigs:
        f.n = sum(1 for g in dfigs if g.page == f.page and g.number <= f.number)  # type: ignore[operator]
    gts["manual"] = build_gt(
        "manual",
        "manual.docx",
        DocKind.docx,
        sha(out / "manual.docx"),
        2,
        dblocks,
        dfigs,
        paged=False,
        header=False,
    )
    for stem, gt in gts.items():
        write_text(gt_dir / f"{stem}.json", dump_json(gt))

    # project-a
    ref = ModelRef(provider="mock", model="mock-1")
    cfg = ProjectConfig(
        name="project-a",
        models=TaskModels(
            cleanup=ref,
            grouping=ref,
            vision=ref,
            embeddings=ref,
            verifier=ModelRef(provider="mock", model="mock-verifier"),
            metadata=ref,
            agent=ref,
        ),
    )
    pa = out / "project-a"
    (pa / "sources").mkdir(parents=True, exist_ok=True)
    write_text(pa / "project.json", dump_json(cfg))
    for name in ("clean.pdf", "figures.pdf"):
        shutil.copyfile(out / name, pa / "sources" / name)

    paths = (
        FIXTURE_FILES
        + [f"ground_truth/{s}.json" for s in STEMS]
        + ["project-a/project.json", "project-a/sources/clean.pdf", "project-a/sources/figures.pdf"]
    )
    sums = {p: sha(out / p) for p in sorted(paths)}
    write_text(out / "CHECKSUMS.json", dump_json(sums))
    return sums


def check(out: Path) -> int:
    expected = out / "CHECKSUMS.json"
    if not expected.exists():
        print(f"missing {expected}", file=sys.stderr)
        return 1
    with tempfile.TemporaryDirectory() as tmp:
        sums = generate(Path(tmp))
        want = __import__("json").loads(expected.read_text())
    bad = sorted(k for k in set(sums) | set(want) if sums.get(k) != want.get(k))
    if bad:
        print("fixture checksum mismatch:", *bad, sep="\n  ", file=sys.stderr)
        return 1
    print(f"ok: {len(sums)} files match CHECKSUMS.json")
    return 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", type=Path, default=Path("tests/fixtures"))
    ap.add_argument("--check", action="store_true")
    args = ap.parse_args(argv)
    if args.check:
        return check(args.out)
    sums = generate(args.out)
    print(f"wrote {len(sums)} files to {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
